"""Generación de reportes en PDF (reportlab) y Excel (openpyxl), todo en memoria.

Seguridad: asuntos y remitentes los escribe el atacante.
  * PDF: se escapan antes de pasar al marcado de reportlab (que interpreta etiquetas como <a> o <img>).
  * Excel: una celda que empieza por «=» se guardaría como fórmula; se fuerza a texto (inyección de fórmulas) y se
    quitan los caracteres de control que Excel no acepta.
"""
from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from centinela.aplicacion.puertos import DatosReporte

ZONA = ZoneInfo("America/Bogota")
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")
NIVEL = {"bajo": "Riesgo bajo", "medio": "Sospechoso", "alto": "Peligroso"}
ESTADO_CASO = {"nuevo": "Por revisar", "en_cuarentena": "En cuarentena", "liberado": "Liberado",
               "confirmado": "Phishing confirmado", "falso_positivo": "Era legítimo"}
ESTADO_BUZON = {"activo": "Protegido", "pausado": "En pausa", "revocado": "Desconectado", "error": "Con problemas"}
INDICADORES = [
    ("suplantacion", "Suplantan a una entidad o empresa"),
    ("enlace", "Enlaces peligrosos (incluye códigos QR)"),
    ("adjunto", "Adjuntos peligrosos"),
    ("fraude", "Fraude sin enlaces (falso jefe o proveedor, extorsión)"),
    ("autenticacion", "Remitente sin autenticar (SPF, DKIM, DMARC)"),
    ("remitente", "Remitente sospechoso"),
    ("contenido", "Texto típico de phishing"),
]
EXTRAS = [("enlaces_qr", "Enlaces escondidos en códigos QR"), ("adjuntos_peligrosos", "Archivos adjuntos peligrosos")]
NOTA_PRIVACIDAD = ("Centinela no guarda el contenido de los correos: este reporte solo incluye los datos del análisis "
                   "(remitente, asunto, nivel de riesgo y razones). Las fechas están en hora de Colombia.")

AZUL = "#1f5fbf"
OSCURO = "#0f1f38"
LOGO = Path(__file__).parent / "recursos" / "logo-blanco.png"
TEXTO = "#142033"
TENUE = "#5b6b82"
BORDE = "#dde4ee"
FONDO = "#f3f6fa"
ALTO = "#b42318"
MEDIO = "#a15c00"
SERIE_MEDIO = "#e3a72f"
SERIE_ALTO = "#a8201a"


def _fecha(d) -> str:
    meses = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
    return f"{d.day} {meses[d.month - 1]} {d.year}"


def _fecha_hora(dt: datetime) -> str:
    local = dt.astimezone(ZONA)
    return f"{_fecha(local)}, {local:%H:%M}"


def _casos_abiertos(d: DatosReporte) -> int:
    return d.estadisticas.casos_por_estado.get("nuevo", 0) + d.estadisticas.casos_por_estado.get("en_cuarentena", 0)


# ====================================================================== PDF
def _seguro(texto: str | None, maximo: int = 300) -> str:
    """Texto del atacante → apto para un Paragraph: recortado, escapado y en caracteres que la fuente puede dibujar."""
    t = " ".join(_CONTROL.sub("", texto or "").split())
    if len(t) > maximo:
        t = t[: maximo - 1] + "…"
    return escape(t.encode("cp1252", "replace").decode("cp1252"))


class GeneradorReportes:
    # ------------------------------------------------------------------ PDF
    def pdf(self, d: DatosReporte) -> bytes:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        c = colors.HexColor
        base = ParagraphStyle("base", fontName="Helvetica", fontSize=9, leading=12, textColor=c(TEXTO))
        p_tenue = ParagraphStyle("tenue", parent=base, textColor=c(TENUE), fontSize=8, leading=10)
        p_celda = ParagraphStyle("celda", parent=base, fontSize=7.5, leading=9.5)
        p_celda_tenue = ParagraphStyle("celda_tenue", parent=p_celda, textColor=c(TENUE))
        h1 = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=18, leading=22, spaceAfter=2)
        h2 = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=11.5, leading=14, spaceBefore=12,
                            spaceAfter=6, textColor=c(OSCURO), keepWithNext=1)
        p_cifra = ParagraphStyle("cifra", parent=base, fontName="Helvetica-Bold", fontSize=18, leading=22)
        ancho = letter[0] - 36 * mm

        def tabla(filas, anchos, encabezado=True, alinear_numeros=True):
            t = Table(filas, colWidths=anchos, repeatRows=1 if encabezado else 0)
            estilo = [("FONT", (0, 0), (-1, -1), "Helvetica", 8.5), ("TEXTCOLOR", (0, 0), (-1, -1), c(TEXTO)),
                      ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.4, c(BORDE)),
                      ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
            if encabezado:
                estilo += [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8), ("TEXTCOLOR", (0, 0), (-1, 0), c(TENUE)),
                           ("BACKGROUND", (0, 0), (-1, 0), c(FONDO))]
            if alinear_numeros:
                estilo.append(("ALIGN", (-1, 0), (-1, -1), "RIGHT"))
            t.setStyle(TableStyle(estilo))
            return t

        e = d.estadisticas
        historia = [
            Paragraph(_seguro(d.titulo), h1),
            Paragraph(f"{_seguro(d.alcance)} · del {_fecha(d.desde)} al {_fecha(d.hasta)} · "
                      f"generado por {_seguro(d.persona, 80)} el {_fecha_hora(d.generado_en)}", p_tenue),
            Spacer(1, 10),
        ]

        # Indicadores principales
        tarjetas = [("Correos revisados", e.revisados, TEXTO), ("Sospechosos", e.por_nivel.get("medio", 0), MEDIO),
                    ("Peligrosos", e.por_nivel.get("alto", 0), ALTO), ("Casos abiertos", _casos_abiertos(d), AZUL)]
        celdas = [[Paragraph(f'<font color="{color}">{n:,}</font>'.replace(",", "."), p_cifra) for _, n, color in tarjetas],
                  [Paragraph(t, p_tenue) for t, _, _ in tarjetas]]
        kpi = Table(celdas, colWidths=[ancho / 4] * 4)
        kpi.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, c(BORDE)), ("INNERGRID", (0, 0), (-1, -1), 0, c("#ffffff")),
                                 ("LINEAFTER", (0, 0), (-2, -1), 0.6, c(BORDE)), ("BACKGROUND", (0, 0), (-1, -1), c("#ffffff")),
                                 ("TOPPADDING", (0, 0), (-1, 0), 8), ("BOTTOMPADDING", (0, 1), (-1, 1), 8),
                                 ("LEFTPADDING", (0, 0), (-1, -1), 10)]))
        historia.append(kpi)

        # Gráfico por día
        historia.append(Paragraph("Correos de riesgo por día", h2))
        if any(x.sospechosos or x.peligrosos for x in e.serie):
            historia.append(self._grafico(d, ancho))
        else:
            historia.append(Paragraph("No hubo correos de riesgo en este periodo.", p_tenue))

        # Tipos de amenaza
        filas = [["Tipo de amenaza detectada", "Correos"]]
        filas += [[t, f"{d.indicadores.get(k, 0):,}".replace(",", ".")] for k, t in INDICADORES]
        filas += [[t, f"{d.indicadores.get(k, 0):,}".replace(",", ".")] for k, t in EXTRAS]
        bloque = [Paragraph("Qué tipo de amenazas llegaron", h2), tabla(filas, [ancho - 30 * mm, 30 * mm]),
                  Paragraph("Un correo puede contar en varios tipos. Se cuentan las razones que se le mostraron a la "
                            "persona (máximo tres por correo).", p_tenue)]
        historia.append(KeepTogether(bloque))

        # Estados y entidades, lado a lado
        estados = [["Estado de los casos", "Casos"]] + [[t, str(e.casos_por_estado.get(k, 0))] for k, t in ESTADO_CASO.items()]
        entidades = [["Entidades más suplantadas", "Correos"]] + (
            [[_seguro(n, 60), str(v)] for n, v in e.entidades_suplantadas] or [["Ninguna en el periodo", "—"]])
        mitad = (ancho - 8 * mm) / 2
        lado = Table([[tabla(estados, [mitad - 22 * mm, 22 * mm]), tabla(entidades, [mitad - 22 * mm, 22 * mm])]],
                     colWidths=[mitad + 4 * mm, mitad + 4 * mm])
        lado.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                  ("RIGHTPADDING", (0, 0), (-1, -1), 8)]))
        historia += [Spacer(1, 6), lado]

        if d.buzones:
            filas = [["Buzón", "Estado", "Revisados", "Sospechosos", "Peligrosos"]] + [
                [_seguro(b.correo, 60), ESTADO_BUZON.get(b.estado, b.estado), str(b.revisados), str(b.sospechosos),
                 str(b.peligrosos)] for b in d.buzones]
            t = tabla(filas, [ancho - 98 * mm, 26 * mm, 24 * mm, 24 * mm, 24 * mm])
            t.setStyle(TableStyle([("ALIGN", (2, 0), (-1, -1), "RIGHT")]))
            historia += [Paragraph("Actividad por buzón", h2), t]
        if d.modelos:
            filas = [["Modelo en uso", "Versión", "Datos de entrenamiento"]] + [
                [f"{m.tipo.capitalize()}: {_seguro(m.nombre, 60)}", _seguro(m.version, 30), _seguro(m.dataset, 60)]
                for m in d.modelos]
            historia += [Paragraph("Modelos de detección", h2),
                         tabla(filas, [ancho * 0.45, ancho * 0.18, ancho * 0.37], alinear_numeros=False)]

        if not d.organizacion:
            historia.append(Paragraph(f"Correos de riesgo del periodo ({len(d.casos)})", h2))
            if d.casos_truncados:
                historia.append(Paragraph(f"Se muestran los {len(d.casos)} más recientes. El Excel trae los mismos.",
                                          p_tenue))
            if not d.casos:
                historia.append(Paragraph("No hubo correos de riesgo en este periodo.", p_tenue))
            else:
                filas = [["Recibido", "Remitente y asunto", "Nivel", "Estado", "Por qué"]]
                for x in d.casos:
                    remitente = x.remitente_nombre or x.remitente_correo or "Remitente desconocido"
                    detalle = (f"<b>{_seguro(remitente, 80)}</b><br/>"
                               f'<font color="{TENUE}">{_seguro(x.remitente_correo, 80)}</font><br/>'
                               f"{_seguro(x.asunto or '(sin asunto)', 160)}")
                    color = ALTO if x.nivel == "alto" else MEDIO
                    razones = "<br/>".join(f"• {_seguro(r, 360)}" for r in x.razones) or "—"
                    filas.append([Paragraph(_fecha_hora(x.recibido_en), p_celda_tenue), Paragraph(detalle, p_celda),
                                  Paragraph(f'<font color="{color}"><b>{NIVEL.get(x.nivel, x.nivel)}</b></font>', p_celda),
                                  Paragraph(ESTADO_CASO.get(x.estado, x.estado), p_celda), Paragraph(razones, p_celda)])
                historia.append(tabla(filas, [22 * mm, 48 * mm, 22 * mm, 21 * mm, ancho - 113 * mm],
                                      alinear_numeros=False))

        historia += [Spacer(1, 14), Paragraph(NOTA_PRIVACIDAD, p_tenue)]

        salida = io.BytesIO()
        doc = SimpleDocTemplate(salida, pagesize=letter, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=24 * mm,
                                bottomMargin=18 * mm, title=d.titulo, author="Centinela", subject=d.alcance,
                                creator="Centinela")
        generado = _fecha_hora(d.generado_en)

        def marco(lienzo, documento):
            lienzo.saveState()
            alto_pagina = documento.pagesize[1]
            lienzo.setFillColor(c(OSCURO))
            lienzo.rect(0, alto_pagina - 12 * mm, documento.pagesize[0], 12 * mm, stroke=0, fill=1)
            lienzo.drawImage(str(LOGO), 18 * mm, alto_pagina - 10.5 * mm, width=8 * mm, height=9 * mm,
                             mask="auto", preserveAspectRatio=True)
            lienzo.setFillColor(c("#ffffff"))
            lienzo.setFont("Helvetica-Bold", 10)
            lienzo.drawString(28 * mm, alto_pagina - 7.6 * mm, "CENTINELA")
            lienzo.setFont("Helvetica", 8)
            lienzo.drawRightString(documento.pagesize[0] - 18 * mm, alto_pagina - 7.6 * mm,
                                   "Protección contra phishing en Gmail")
            lienzo.setFillColor(c(TENUE))
            lienzo.drawString(18 * mm, 10 * mm, f"Generado el {generado}")
            lienzo.drawRightString(documento.pagesize[0] - 18 * mm, 10 * mm, f"Página {documento.page}")
            lienzo.restoreState()

        doc.build(historia, onFirstPage=marco, onLaterPages=marco)
        return salida.getvalue()

    @staticmethod
    def _grafico(d: DatosReporte, ancho: float):
        from reportlab.graphics.charts.barcharts import VerticalBarChart
        from reportlab.graphics.charts.legends import Legend
        from reportlab.graphics.shapes import Drawing
        from reportlab.lib import colors

        serie = d.estadisticas.serie
        alto = 150
        dibujo = Drawing(ancho, alto + 20)
        g = VerticalBarChart()
        g.x, g.y, g.width, g.height = 28, 22, ancho - 40, alto - 20
        g.data = [[x.sospechosos for x in serie], [x.peligrosos for x in serie]]
        g.categoryAxis.style = "stacked"
        paso = max(1, len(serie) // 10)
        g.categoryAxis.categoryNames = [f"{x.fecha.day}/{x.fecha.month}" if i % paso == 0 else "" for i, x in enumerate(serie)]
        g.categoryAxis.labels.fontName = "Helvetica"
        g.categoryAxis.labels.fontSize = 6.5
        g.categoryAxis.labels.fillColor = colors.HexColor(TENUE)
        g.categoryAxis.strokeColor = colors.HexColor(BORDE)
        g.valueAxis.valueMin = 0
        maximo = max((x.sospechosos + x.peligrosos for x in serie), default=1)
        g.valueAxis.valueStep = max(1, round(maximo / 4))
        g.valueAxis.labels.fontName = "Helvetica"
        g.valueAxis.labels.fontSize = 6.5
        g.valueAxis.labels.fillColor = colors.HexColor(TENUE)
        g.valueAxis.strokeColor = colors.HexColor(BORDE)
        g.valueAxis.visibleGrid = True
        g.valueAxis.gridStrokeColor = colors.HexColor("#eef2f7")
        g.groupSpacing = 1 if len(serie) > 40 else 3
        g.bars.strokeColor = None
        g.bars[0].fillColor = colors.HexColor(SERIE_MEDIO)
        g.bars[1].fillColor = colors.HexColor(SERIE_ALTO)
        dibujo.add(g)
        ley = Legend()
        ley.x, ley.y = 30, alto + 12
        ley.alignment = "right"
        ley.columnMaximum = 1
        ley.deltax = 90
        ley.fontName = "Helvetica"
        ley.fontSize = 7.5
        ley.fillColor = colors.HexColor(TEXTO)
        ley.boxAnchor = "w"
        ley.colorNamePairs = [(colors.HexColor(SERIE_MEDIO), "Sospechosos"), (colors.HexColor(SERIE_ALTO), "Peligrosos")]
        dibujo.add(ley)
        return dibujo

    # ------------------------------------------------------------------ Excel
    def excel(self, d: DatosReporte) -> bytes:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        libro = Workbook()
        titulo = Font(bold=True, size=14, color=OSCURO[1:])
        tenue = Font(color=TENUE[1:], size=9)
        relleno = PatternFill("solid", fgColor=FONDO[1:])
        encabezado = Font(bold=True, color=OSCURO[1:])

        def fila(hoja, valores, fuente=None):
            hoja.append([_celda(v) for v in valores])
            for celda in hoja[hoja.max_row]:
                if isinstance(celda.value, str):
                    celda.data_type = "s"           # nunca una fórmula, aunque el texto empiece por «=»
                if fuente:
                    celda.font = fuente

        def cabecera(hoja, valores):
            fila(hoja, valores, encabezado)
            for celda in hoja[hoja.max_row]:
                celda.fill = relleno

        def anchos(hoja, valores):
            for i, w in enumerate(valores, 1):
                hoja.column_dimensions[get_column_letter(i)].width = w

        e = d.estadisticas
        # --- Resumen
        r = libro.active
        r.title = "Resumen"
        fila(r, [d.titulo], titulo)
        fila(r, [f"{d.alcance} · del {_fecha(d.desde)} al {_fecha(d.hasta)}"], tenue)
        fila(r, [f"Generado por {d.persona} el {_fecha_hora(d.generado_en)}"], tenue)
        fila(r, [])
        cabecera(r, ["Indicador", "Valor"])
        for t, v in [("Correos revisados", e.revisados), ("Sospechosos", e.por_nivel.get("medio", 0)),
                     ("Peligrosos", e.por_nivel.get("alto", 0)), ("Casos abiertos", _casos_abiertos(d))]:
            fila(r, [t, v])
        fila(r, [])
        cabecera(r, ["Tipo de amenaza detectada", "Correos"])
        for k, t in INDICADORES + EXTRAS:
            fila(r, [t, d.indicadores.get(k, 0)])
        fila(r, [])
        cabecera(r, ["Estado de los casos", "Casos"])
        for k, t in ESTADO_CASO.items():
            fila(r, [t, e.casos_por_estado.get(k, 0)])
        fila(r, [])
        cabecera(r, ["Entidades más suplantadas", "Correos"])
        for n, v in e.entidades_suplantadas or [("Ninguna en el periodo", 0)]:
            fila(r, [n, v])
        fila(r, [])
        fila(r, [NOTA_PRIVACIDAD], tenue)
        anchos(r, [58, 14])

        # --- Por día (con gráfico nativo de Excel)
        dia = libro.create_sheet("Por día")
        cabecera(dia, ["Fecha", "Revisados", "Sospechosos", "Peligrosos"])
        for x in e.serie:
            fila(dia, [x.fecha, x.revisados, x.sospechosos, x.peligrosos])
            dia.cell(dia.max_row, 1).number_format = "dd/mm/yyyy"
        anchos(dia, [13, 12, 13, 12])
        dia.freeze_panes = "A2"
        if e.serie:
            grafico = BarChart()
            grafico.type, grafico.grouping, grafico.overlap = "col", "stacked", 100
            grafico.title = "Correos de riesgo por día"
            grafico.y_axis.title = "Correos"
            grafico.height, grafico.width = 8, 22
            grafico.add_data(Reference(dia, min_col=3, max_col=4, min_row=1, max_row=dia.max_row), titles_from_data=True)
            grafico.set_categories(Reference(dia, min_col=1, min_row=2, max_row=dia.max_row))
            grafico.series[0].graphicalProperties.solidFill = SERIE_MEDIO[1:]
            grafico.series[1].graphicalProperties.solidFill = SERIE_ALTO[1:]
            dia.add_chart(grafico, "F2")

        # --- Correos de riesgo (solo el reporte personal)
        if not d.organizacion:
            cs = libro.create_sheet("Correos de riesgo")
            cabecera(cs, ["Caso", "Recibido", "Buzón", "Remitente", "Correo del remitente", "Asunto", "Nivel",
                          "Riesgo estimado", "Estado", "Razón 1", "Razón 2", "Razón 3"])
            for x in d.casos:
                razones = (list(x.razones) + ["", "", ""])[:3]
                fila(cs, [x.id, x.recibido_en.astimezone(ZONA).replace(tzinfo=None), x.buzon_correo, x.remitente_nombre or "",
                          x.remitente_correo or "", x.asunto or "(sin asunto)", NIVEL.get(x.nivel, x.nivel), x.p_final,
                          ESTADO_CASO.get(x.estado, x.estado), *razones])
                cs.cell(cs.max_row, 2).number_format = "dd/mm/yyyy hh:mm"
                cs.cell(cs.max_row, 8).number_format = "0%"
                for col in (10, 11, 12):
                    cs.cell(cs.max_row, col).alignment = Alignment(wrap_text=True, vertical="top")
            anchos(cs, [7, 17, 26, 24, 28, 44, 12, 10, 18, 60, 60, 60])
            cs.freeze_panes = "C2"
            cs.auto_filter.ref = f"A1:L{max(cs.max_row, 1)}"

        # --- Por buzón y modelos (reporte de la organización)
        if d.buzones:
            bz = libro.create_sheet("Por buzón")
            cabecera(bz, ["Buzón", "Estado", "Revisados", "Sospechosos", "Peligrosos"])
            for b in d.buzones:
                fila(bz, [b.correo, ESTADO_BUZON.get(b.estado, b.estado), b.revisados, b.sospechosos, b.peligrosos])
            anchos(bz, [36, 15, 12, 13, 12])
            bz.freeze_panes = "A2"
            bz.auto_filter.ref = f"A1:E{bz.max_row}"
        if d.modelos:
            md = libro.create_sheet("Modelos")
            cabecera(md, ["Tipo", "Nombre", "Versión", "Datos de entrenamiento", "Umbral medio", "Umbral alto"])
            for m in d.modelos:
                fila(md, [m.tipo, m.nombre, m.version, m.dataset, m.umbral_medio, m.umbral_alto])
            anchos(md, [12, 34, 16, 40, 13, 12])

        for hoja in libro.worksheets:
            hoja.sheet_view.showGridLines = hoja.title != "Resumen"
        libro.properties.title = d.titulo
        libro.properties.creator = "Centinela"
        salida = io.BytesIO()
        libro.save(salida)
        return salida.getvalue()


def _celda(v):
    """Quita los caracteres de control que Excel no acepta (vienen en asuntos maliciosos) y recorta textos largos."""
    if isinstance(v, str):
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

        return ILLEGAL_CHARACTERS_RE.sub("", v)[:32000]
    return v
