"""Avisos de Gmail en tiempo real (token de Google, users.watch, worker) y generación de reportes PDF y Excel."""
from __future__ import annotations

import io
import json
import time
from datetime import UTC, date, datetime, timedelta

import httpx
import jwt
import pytest
from cryptography.hazmat.primitives.asymmetric import rsa
from openpyxl import load_workbook

from centinela.adaptadores.salida.gmail import ProveedorGmail
from centinela.adaptadores.salida.lector_correo import LectorCorreoMime
from centinela.adaptadores.salida.pubsub import VerificadorOidcGoogle
from centinela.adaptadores.salida.reportes import GeneradorReportes
from centinela.aplicacion.casos_uso import buzones as bz
from centinela.aplicacion.casos_uso import tiempo_real as tr
from centinela.aplicacion.puertos import (
    AvisoNoAutentico,
    DatosReporte,
    DiaPanel,
    ErrorProveedorCorreo,
    EstadisticasPanel,
    FilaBuzonReporte,
    ModeloRegistrado,
    ResumenCaso,
)
from centinela.dominio.errores import CredencialesInvalidas, ServicioNoDisponible
from tests.fakes import RelojFijo
from tests.gmail_falso import correo
from tests.unit.test_casos_uso_buzones import conectar
from tests.unit.test_casos_uso_buzones import s as s  # noqa: F401 - fixture compartida
from tests.unit.test_gmail_adaptador import GoogleSimulado, token_ok

AUDIENCIA = "https://centinela.example.co/api/gmail/notificaciones"
CUENTA = "pubsub-push@centinela-tesis.iam.gserviceaccount.com"
TEMA = "projects/centinela-tesis/topics/gmail"


# ---------------------------------------------------------------------- token OIDC de Pub/Sub
@pytest.fixture(scope="module")
def llaves():
    privada = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    jwk = json.loads(jwt.algorithms.RSAAlgorithm.to_jwk(privada.public_key()))
    jwk.update(kid="k1", alg="RS256", use="sig")
    return privada, jwk


def _verificador(jwk) -> tuple[VerificadorOidcGoogle, list]:
    descargas = []

    def google(req):
        descargas.append(req)
        return httpx.Response(200, json={"keys": [jwk]})

    return VerificadorOidcGoogle(AUDIENCIA, CUENTA, httpx.Client(transport=httpx.MockTransport(google))), descargas


def _token(privada, kid="k1", alg="RS256", **cambios) -> str:
    ahora = int(time.time())
    reclamos = {"iss": "https://accounts.google.com", "aud": AUDIENCIA, "email": CUENTA, "email_verified": True,
                "iat": ahora, "exp": ahora + 3600, "sub": "1234"}
    reclamos.update(cambios)
    return jwt.encode({k: v for k, v in reclamos.items() if v is not None}, privada, algorithm=alg, headers={"kid": kid})


def test_token_de_google_valido(llaves):
    privada, jwk = llaves
    v, descargas = _verificador(jwk)
    v.verificar(_token(privada))
    v.verificar(_token(privada))
    assert len(descargas) == 1                                # las llaves de Google se guardan en memoria


@pytest.mark.parametrize("cambios", [
    {"aud": "https://otro-sitio.co"}, {"email": "atacante@evil.iam.gserviceaccount.com"}, {"email_verified": False},
    {"iss": "https://evil.co"}, {"exp": int(time.time()) - 600}, {"kid": "desconocida"}, {"exp": None}])
def test_tokens_rechazados(llaves, cambios):
    privada, jwk = llaves
    v, _ = _verificador(jwk)
    kid = cambios.pop("kid", "k1")
    with pytest.raises(AvisoNoAutentico):
        v.verificar(_token(privada, kid=kid, **cambios))


def test_token_firmado_por_otro_o_mal_formado(llaves):
    _, jwk = llaves
    v, _ = _verificador(jwk)
    otra = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    for token in (_token(otra), "no-es-un-jwt", jwt.encode({"aud": AUDIENCIA}, "x" * 32, algorithm="HS256",
                                                          headers={"kid": "k1"})):
        with pytest.raises(AvisoNoAutentico):
            v.verificar(token)


# ---------------------------------------------------------------------- users.watch
def test_vigilar_pide_la_bandeja_y_un_403_no_desconecta_el_buzon():
    g = GoogleSimulado()
    g.rutas[("POST", "/token")] = token_ok
    g.rutas[("POST", "/gmail/v1/users/me/watch")] = lambda r: httpx.Response(
        200, json={"historyId": "5501", "expiration": "1793000000000"})
    v = ProveedorGmail("cid", "s", g.cliente(), RelojFijo()).abrir("rt").vigilar(TEMA)
    assert v.history_id == 5501 and v.expira_en == datetime.fromtimestamp(1793000000, UTC)
    assert json.loads(g.peticiones[-1].content) == {"topicName": TEMA, "labelIds": ["INBOX"],
                                                    "labelFilterBehavior": "INCLUDE"}
    g.rutas[("POST", "/gmail/v1/users/me/watch")] = lambda r: httpx.Response(403, json={"error": {
        "message": "User not authorized to perform this action.", "errors": [{"reason": "forbidden"}]}})
    with pytest.raises(ErrorProveedorCorreo, match="not authorized"):
        ProveedorGmail("cid", "s", g.cliente(), RelojFijo()).abrir("rt").vigilar(TEMA)


# ---------------------------------------------------------------------- casos de uso
class VerificadorFalso:
    def verificar(self, token):
        if token != "bueno":
            raise AvisoNoAutentico("firma inválida")


def test_aviso_de_gmail_pide_revisar_ese_buzon(s):
    b = conectar(s)
    recibir = tr.RecibirAvisoGmail(s.uow, s.reloj, VerificadorFalso())
    assert recibir.ejecutar("bueno", "ANA@gmail.com") is True
    assert s.uow.almacen.revision_solicitada == {b.id}
    assert recibir.ejecutar("bueno", "desconocido@gmail.com") is False
    assert recibir.ejecutar("bueno", None) is False
    with pytest.raises(CredencialesInvalidas):
        recibir.ejecutar("falso", "ana@gmail.com")
    with pytest.raises(CredencialesInvalidas):
        recibir.ejecutar(None, "ana@gmail.com")
    with pytest.raises(ServicioNoDisponible):
        tr.RecibirAvisoGmail(s.uow, s.reloj, None).ejecutar("bueno", "ana@gmail.com")


def test_aviso_de_un_buzon_en_pausa_se_ignora(s):
    b = conectar(s)
    s.actualizar.ejecutar(2, b.id, pausado=True)
    assert tr.RecibirAvisoGmail(s.uow, s.reloj, VerificadorFalso()).ejecutar("bueno", "ana@gmail.com") is False


def test_el_worker_revisa_al_instante_lo_que_gmail_aviso(s):
    b = conectar(s)
    pendientes = bz.RevisarBuzonesPendientes(s.uow, s.reloj, s.revisar, tiempo_real=True)
    s.uow.almacen.buzones[b.id].vigilancia_expira_en = s.reloj.ahora() + timedelta(days=7)
    assert [r.buzon_id for r in pendientes.ejecutar()] == [b.id]          # primera revisión
    s.reloj.avanzar(seconds=40)
    s.gmail.llega(correo())
    assert pendientes.ejecutar() == []          # vigilado: sin aviso no se revisa a los 30 s (el respaldo es de 5 min)
    tr.RecibirAvisoGmail(s.uow, s.reloj, VerificadorFalso()).ejecutar("bueno", "ana@gmail.com")
    assert [r.nuevos for r in pendientes.ejecutar()] == [1]               # el aviso lo trae de inmediato
    s.reloj.avanzar(seconds=301)
    assert [r.buzon_id for r in pendientes.ejecutar()] == [b.id]          # revisión de respaldo


def test_renovar_vigilancias(s):
    b = conectar(s)
    renovar = tr.RenovarVigilancias(s.uow, s.reloj, s.proveedor, s.cif, TEMA)
    assert renovar.activo and renovar.ejecutar() == 1
    assert s.gmail.vigilado_en == TEMA
    assert s.uow.almacen.buzones[b.id].en_tiempo_real(datetime.now(UTC))
    assert renovar.ejecutar() == 0                                        # vence en 7 días: aún no toca
    assert not tr.RenovarVigilancias(s.uow, s.reloj, s.proveedor, s.cif, None).activo


def test_si_gmail_rechaza_la_vigilancia_espera_antes_de_reintentar(s):
    conectar(s)
    s.gmail.falla = "tema sin permiso"
    renovar = tr.RenovarVigilancias(s.uow, s.reloj, s.proveedor, s.cif, TEMA)
    assert renovar.ejecutar() == 0
    s.gmail.falla = None
    assert renovar.ejecutar() == 0                         # sigue esperando: no insiste en cada ciclo
    s.reloj.avanzar(minutes=16)
    assert renovar.ejecutar() == 1


def test_reconectar_o_desconectar_borra_la_vigilancia(s):
    b = conectar(s)
    tr.RenovarVigilancias(s.uow, s.reloj, s.proveedor, s.cif, TEMA).ejecutar()
    s.desconectar.ejecutar(2, b.id)
    assert s.uow.almacen.buzones[b.id].vigilancia_expira_en is None


# ---------------------------------------------------------------------- reportes
MALICIOSO = '=HYPERLINK("http://evil.top","Clic") <b>negrita</b> <img src=x> 🚨\x07 & «Bancolombia»'


def _datos(organizacion=False, casos=2) -> DatosReporte:
    hoy = date(2026, 10, 20)
    serie = [DiaPanel(hoy - timedelta(days=29 - i), 10 + i, i % 3, i % 2) for i in range(30)]
    e = EstadisticasPanel(revisados=sum(x.revisados for x in serie), por_nivel={"bajo": 300, "medio": 29, "alto": 15},
                          casos_por_estado={"nuevo": 3, "en_cuarentena": 5, "liberado": 1, "confirmado": 4,
                                            "falso_positivo": 1},
                          entidades_suplantadas=[("Bancolombia", 6), ("DIAN", 2)], serie=serie)
    resumenes = [ResumenCaso(id=i, estado="en_cuarentena", creado_en=datetime(2026, 10, 19, 15, tzinfo=UTC),
                             buzon_correo="ana@gmail.com", remitente_nombre=MALICIOSO,
                             remitente_correo="alertas@bancolombia-seguro.xyz", asunto=MALICIOSO,
                             recibido_en=datetime(2026, 10, 19, 15, 30, tzinfo=UTC), nivel="alto", p_final=0.93,
                             resumen="r", razones=["Suplanta a Bancolombia", "Un enlace lleva a evil.top", MALICIOSO])
                 for i in range(1, casos + 1)]
    return DatosReporte(
        titulo="Reporte de seguridad del correo", alcance="Toda la organización" if organizacion else "Mis buzones",
        persona="Ana <script>", generado_en=datetime(2026, 10, 20, 17, 5, tzinfo=UTC), desde=serie[0].fecha,
        hasta=hoy, estadisticas=e, organizacion=organizacion,
        indicadores={"suplantacion": 8, "enlace": 12, "adjunto": 2, "fraude": 3, "autenticacion": 6, "remitente": 1,
                     "contenido": 9, "enlaces_qr": 2, "adjuntos_peligrosos": 3},
        casos=[] if organizacion else resumenes,
        buzones=[FilaBuzonReporte("ana@gmail.com", "activo", 300, 20, 10)] if organizacion else [],
        modelos=[ModeloRegistrado(1, "contenido", "Texto", "v2", "r", "0" * 64, "corpus", 0.3, 0.7, True, {})]
        if organizacion else [])


@pytest.mark.parametrize("organizacion", [False, True])
def test_pdf_se_genera_con_textos_del_atacante(organizacion):
    pdf = GeneradorReportes().pdf(_datos(organizacion, casos=60))
    assert pdf.startswith(b"%PDF-") and pdf.rstrip().endswith(b"%%EOF") and pdf.count(b"/Type /Page\n") >= 2 - organizacion


def test_pdf_sin_actividad():
    d = _datos()
    vacio = DatosReporte(**{**d.__dict__, "casos": [], "estadisticas": EstadisticasPanel(
        0, {}, {}, [], [DiaPanel(d.desde + timedelta(days=i), 0, 0, 0) for i in range(30)])})
    assert GeneradorReportes().pdf(vacio).startswith(b"%PDF-")


def test_excel_personal_sin_formulas_inyectadas():
    libro = load_workbook(io.BytesIO(GeneradorReportes().excel(_datos())))
    assert libro.sheetnames == ["Resumen", "Por día", "Correos de riesgo"]
    hoja = libro["Correos de riesgo"]
    encabezados = [c.value for c in hoja[1]]
    assert encabezados[:6] == ["Caso", "Recibido", "Buzón", "Remitente", "Correo del remitente", "Asunto"]
    asunto = hoja.cell(2, 6)
    assert asunto.data_type == "s" and asunto.value.startswith("=HYPERLINK") and "\x07" not in asunto.value
    assert hoja.cell(2, 2).value == datetime(2026, 10, 19, 10, 30)                 # hora de Colombia
    assert hoja.auto_filter.ref == "A1:L3" and len(libro["Por día"]._charts) == 1
    for fila in libro["Resumen"].iter_rows():
        for celda in fila:
            assert celda.data_type != "f"


def test_excel_de_la_organizacion_no_trae_asuntos():
    libro = load_workbook(io.BytesIO(GeneradorReportes().excel(_datos(organizacion=True))))
    assert libro.sheetnames == ["Resumen", "Por día", "Por buzón", "Modelos"]
    textos = " ".join(str(c.value) for h in libro.worksheets for fila in h.iter_rows() for c in fila if c.value)
    assert "HYPERLINK" not in textos and "ana@gmail.com" in textos


def test_lector_sin_inspector_sigue_funcionando():
    c = LectorCorreoMime(b"k").leer(correo(adjuntos=[("a.zip", b"PK\x03\x04", "application/zip")]))
    assert c.adjuntos[0].senales == () and c.adjuntos[0].tipo_real is None
