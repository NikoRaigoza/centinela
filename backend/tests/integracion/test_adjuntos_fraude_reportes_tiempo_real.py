"""Funciones nuevas contra PostgreSQL real: adjuntos y códigos QR, fraude sin enlaces, reportes PDF/Excel y avisos de
Gmail en tiempo real (endpoint de Pub/Sub y despertador del worker)."""
from __future__ import annotations

import base64
import io
import json
import threading
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from centinela.adaptadores.salida.postgres.despertador import Despertador
from centinela.aplicacion.puertos import AvisoNoAutentico
from tests.gmail_falso import correo
from tests.integracion.test_flujo_analisis import LEGITIMO, armar_entorno, cliente_de, conectar_buzon
from tests.unit.test_adjuntos_qr_y_fraude import qr_png, un_zip

pytestmark = pytest.mark.integracion
AUTH = "mx.google.com; dkim=pass; spf=pass; dmarc=pass"

QUISHING = correo(remitente="Seguridad Microsoft 365 <no-reply@ms365-secure.com>", asunto="Reactive su verificación",
                  cuerpo="Su método de verificación vence hoy. Escanee el código QR con su celular para no perder el acceso.",
                  autenticacion=AUTH, message_id="<qr@x>",
                  adjuntos=[("codigo.png", qr_png("https://ms365-login.top/auth?u=ana"), "image/png"),
                            ("factura.zip", un_zip({"Factura.pdf.exe": b"MZ" + b"\x00" * 64}), "application/zip")])
FRAUDE = correo(remitente="Carlos Ramírez Gerente <gerencia.carlos@gmail.com>", asunto="Favor",
                cuerpo="Hola, ¿estás disponible? Estoy en una reunión y no puedo atender llamadas. Necesito que compres "
                       "5 tarjetas de regalo de Google Play y me envíes los códigos lo antes posible. Es confidencial.",
                autenticacion=AUTH, message_id="<bec@x>")


class VerificadorFalso:
    def verificar(self, token):
        if token != "token-de-google":
            raise AvisoNoAutentico("firma inválida")


@pytest.fixture
def entorno(config, limpiar, motor_owner):
    c, auth, gmail = armar_entorno(config, motor_owner, verificador_gmail=VerificadorFalso())
    yield c, auth, gmail
    c.motor.dispose()


def cliente_admin(c):
    """El administrador entra con doble factor (lo prueba test_api_acceso); aquí se le emite el token directamente."""
    from fastapi.testclient import TestClient

    from centinela.adaptadores.entrada.api.app import crear_app

    with c.uow() as uow:
        admin = uow.usuarios.por_correo("admin@centinela.co")
    cli = TestClient(crear_app(c), base_url="http://testserver")
    cli.headers["Authorization"] = f"Bearer {c.tokens.acceso(admin)}"
    return cli


def _con_casos(entorno):
    c, auth, gmail = entorno
    ana = cliente_de(c, "ana@gmail.com")
    buzon_id = conectar_buzon(ana, auth, gmail)
    for crudo in (QUISHING, FRAUDE, LEGITIMO):
        gmail.llega(crudo)
    assert ana.post(f"/api/buzones/{buzon_id}/revision").json()["nuevos"] == 3
    return c, ana, buzon_id


def test_qr_adjuntos_y_fraude_llegan_al_detalle_del_caso(entorno, motor_owner):
    c, ana, _ = _con_casos(entorno)
    casos = {x["asunto"]: x for x in ana.get("/api/casos").json()}
    assert set(casos) == {"Reactive su verificación", "Favor"}                     # el legítimo no abre caso

    qr = ana.get(f"/api/casos/{casos['Reactive su verificación']['id']}").json()
    assert qr["nivel"] == "alto" and qr["estado"] == "en_cuarentena"
    assert (qr["remitente_nombre"], qr["remitente_correo"]) == ("Seguridad Microsoft 365", "no-reply@ms365-secure.com")
    assert {"dominio": "ms365-login.top", "origen": "qr"}.items() <= next(
        e for e in qr["enlaces"] if e["origen"] == "qr").items()
    adjuntos = {a["nombre"]: a for a in qr["adjuntos"]}
    assert adjuntos["codigo.png"]["peligroso"] is False
    assert adjuntos["factura.zip"]["peligroso"] is True and "Factura.pdf.exe" in adjuntos["factura.zip"]["motivo"]
    assert any("código QR" in r for r in qr["razones"])

    bec = ana.get(f"/api/casos/{casos['Favor']['id']}").json()
    assert bec["nivel"] == "alto" and bec["enlaces"] == [] and bec["adjuntos"] == []
    assert bec["todas_las_razones"][0]["area"] == "fraude"
    assert "tarjetas de regalo" in bec["razones"][0]

    with motor_owner.connect() as cx:          # lo guardado en la BD: origen del enlace y motivo del adjunto
        assert cx.scalar(text("SELECT count(*) FROM centinela.enlace_analizado WHERE origen = 'qr'")) == 1
        assert cx.scalar(text("SELECT count(*) FROM centinela.razon WHERE area = 'fraude'")) == 1
        assert cx.scalar(text("SELECT count(*) FROM centinela.adjunto_analizado WHERE peligroso")) == 1


def test_analisis_manual_muestra_adjuntos_y_origen(entorno):
    c, *_ = entorno
    ana = cliente_de(c, "ana@gmail.com")
    r = ana.post("/api/analisis", content=QUISHING, headers={"Content-Type": "message/rfc822"}).json()
    assert r["nivel"] == "alto" and [a["peligroso"] for a in r["adjuntos"]] == [False, True]
    assert r["enlaces"][0]["origen"] == "qr"


def test_reporte_personal_pdf_y_excel(entorno, motor_owner):
    c, ana, _ = _con_casos(entorno)
    pdf = ana.get("/api/reportes", params={"formato": "pdf", "dias": 30})
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF-")
    assert pdf.headers["content-disposition"].startswith('attachment; filename="centinela-reporte-personal-')

    xlsx = ana.get("/api/reportes", params={"formato": "xlsx", "dias": 7})
    assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats")
    libro = load_workbook(io.BytesIO(xlsx.content))
    asuntos = {libro["Correos de riesgo"].cell(f, 6).value for f in (2, 3)}
    assert asuntos == {"Reactive su verificación", "Favor"}
    assert len(list(libro["Por día"].iter_rows())) == 8                               # encabezado + 7 días

    d = c.generar_reporte.datos(2, 30)                      # cifras que alimentan ambos formatos
    assert d.estadisticas.revisados == 3 and d.estadisticas.por_nivel["alto"] == 2
    assert d.indicadores["fraude"] == 1 and d.indicadores["enlaces_qr"] == 1 and d.indicadores["adjuntos_peligrosos"] == 1

    luis = cliente_de(c, "luis@gmail.com")                  # otra persona: su reporte no trae los casos de Ana
    vacio = load_workbook(io.BytesIO(luis.get("/api/reportes", params={"formato": "xlsx"}).content))
    assert vacio["Correos de riesgo"].max_row == 1

    with motor_owner.connect() as cx:
        assert cx.scalar(text("SELECT count(*) FROM centinela.auditoria WHERE accion = 'reporte.descargar'")) == 3


def test_reporte_de_la_organizacion_solo_para_administradores(entorno):
    c, ana, _ = _con_casos(entorno)
    assert ana.get("/api/reportes", params={"organizacion": True}).status_code == 403
    admin = cliente_admin(c)
    r = admin.get("/api/reportes", params={"formato": "xlsx", "organizacion": True})
    assert r.status_code == 200 and "organizacion" in r.headers["content-disposition"]
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames == ["Resumen", "Por día", "Por buzón", "Modelos"]
    assert [c.value for c in libro["Por buzón"][2]] == ["ana@gmail.com", "Protegido", 3, 0, 2]
    textos = " ".join(str(x.value) for h in libro.worksheets for f in h.iter_rows() for x in f if x.value)
    assert "Reactive su verificación" not in textos and "gerencia.carlos" not in textos       # sin asuntos ni remitentes
    assert admin.get("/api/reportes", params={"formato": "pdf", "organizacion": True}).content.startswith(b"%PDF-")
    assert admin.get("/api/reportes", params={"dias": 0}).status_code == 422
    assert admin.get("/api/reportes", params={"formato": "docx"}).status_code == 422


def _sobre(correo_gmail: str) -> bytes:
    datos = base64.b64encode(json.dumps({"emailAddress": correo_gmail, "historyId": 9999}).encode()).decode()
    return json.dumps({"message": {"data": datos, "messageId": "1", "publishTime": "2026-10-20T15:00:00Z"},
                       "subscription": "projects/centinela-tesis/subscriptions/gmail-push"}).encode()


def test_notificacion_de_gmail_revisa_el_buzon_al_instante(entorno, motor_owner):
    c, auth, gmail = entorno
    ana = cliente_de(c, "ana@gmail.com")
    buzon_id = conectar_buzon(ana, auth, gmail)
    with motor_owner.begin() as cx:                         # Gmail ya avisa en tiempo real a este buzón
        cx.execute(text("UPDATE centinela.buzon SET vigilancia_expira_en = now() + interval '7 days', "
                        "ultima_revision_en = now()"))
    assert ana.get("/api/buzones").json()[0]["tiempo_real"] is True
    gmail.llega(FRAUDE)
    assert c.revisar_pendientes.ejecutar() == []            # vigilado: sin aviso, no se revisa todavía

    sin_sesion = cliente_de(c, "luis@gmail.com")
    sin_sesion.headers.pop("Authorization")
    url = "/api/gmail/notificaciones"
    assert sin_sesion.post(url, content=_sobre("ana@gmail.com")).status_code == 401
    assert sin_sesion.post(url, content=_sobre("ana@gmail.com"),
                           headers={"Authorization": "Bearer token-falso"}).status_code == 401
    google = {"Authorization": "Bearer token-de-google"}
    assert sin_sesion.post(url, content=b"basura", headers=google).status_code == 204          # se descarta sin reintentos
    assert sin_sesion.post(url, content=_sobre("otro@gmail.com"), headers=google).status_code == 204
    with motor_owner.connect() as cx:
        assert cx.scalar(text("SELECT count(*) FROM centinela.buzon WHERE revision_solicitada_en IS NOT NULL")) == 0

    assert sin_sesion.post(url, content=_sobre("ana@gmail.com"), headers=google).status_code == 204
    assert [(r.buzon_id, r.nuevos) for r in c.revisar_pendientes.ejecutar()] == [(buzon_id, 1)]
    with motor_owner.connect() as cx:
        assert cx.scalar(text("SELECT revision_solicitada_en FROM centinela.buzon")) is None
    assert c.revisar_pendientes.ejecutar() == []


def test_sin_tiempo_real_configurado_el_endpoint_no_existe_para_nadie(config, limpiar, motor_owner):
    c, *_ = armar_entorno(config, motor_owner)
    cli = cliente_de(c, "ana@gmail.com")
    assert cli.post("/api/gmail/notificaciones", content=_sobre("ana@gmail.com"),
                    headers={"Authorization": "Bearer x"}).status_code == 503
    c.motor.dispose()


def test_despertador_oye_el_aviso_de_la_api(entorno, config):
    c, auth, gmail = entorno
    ana = cliente_de(c, "ana@gmail.com")
    buzon_id = conectar_buzon(ana, auth, gmail)
    despertar, parar = threading.Event(), threading.Event()
    d = Despertador(config.bd_url, despertar, parar).iniciar()
    try:
        assert d.escuchando.wait(10)
        with c.uow() as uow:
            uow.buzones.solicitar_revision(buzon_id, c.reloj.ahora())
            assert not despertar.wait(1.5)                  # NOTIFY solo se entrega al confirmar la transacción
            uow.confirmar()
        assert despertar.wait(5)
    finally:
        parar.set()


def test_remitente_conocido_y_buzones_por_vigilar(entorno):
    c, auth, gmail = entorno
    ana = cliente_de(c, "ana@gmail.com")
    buzon_id = conectar_buzon(ana, auth, gmail)
    gmail.llega(LEGITIMO)
    ana.post(f"/api/buzones/{buzon_id}/revision")
    ahora = c.reloj.ahora()
    with c.uow() as uow:
        assert uow.mensajes.remitente_conocido(buzon_id, "Laura.Gomez@gmail.com")
        assert not uow.mensajes.remitente_conocido(buzon_id, "otro@gmail.com")
        assert [b.id for b in uow.buzones.por_vigilar(ahora + timedelta(days=6), 10)] == [buzon_id]
        b = uow.buzones.por_id(buzon_id)
        b.registrar_vigilancia(ahora + timedelta(days=7))
        uow.buzones.guardar(b)
        uow.confirmar()
    with c.uow() as uow:
        assert uow.buzones.por_vigilar(ahora + timedelta(days=6), 10) == []
        assert uow.buzones.por_id(buzon_id).en_tiempo_real(ahora)
